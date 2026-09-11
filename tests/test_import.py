import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from aml_reconcile import DEFAULT_CONFIG, process, read_file, normalize


class ImportTests(unittest.TestCase):
    def test_negative_completion_and_result_statuses_are_not_successes(self):
        frame = pd.DataFrame([{
            "User ID": "001", "Email": "learner@example.org", "Full Name": "Example",
            "Assignment ID": "A-1", "Assignment Title": "AML",
            "Completion Status": "Not Completed", "Result Status": "Not Passed",
        }])
        records = normalize(frame, Path("input.csv"), "CSV", DEFAULT_CONFIG, 0)
        self.assertEqual(int(records.loc[0, "Completed Flag"]), 0)
        self.assertEqual(int(records.loc[0, "Passed Flag"]), 0)
        self.assertEqual(int(records.loc[0, "Started Not Completed Flag"]), 0)
        self.assertEqual(records.loc[0, "Status Review"], "Unknown completion status")
        self.assertTrue(records.loc[0, "Missing Critical Fields"] == "")
    def test_excel_title_and_blank_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / 'training.xlsx'
            pd.DataFrame([
                ['Training report', None, None],
                [None, None, None],
                ['User ID', 'Email', 'Assignment Title'],
                ['001', 'learner@example.org', 'AML'],
            ]).to_excel(source, index=False, header=False)
            frame = read_file(source)[0][1]
            self.assertEqual(frame.iloc[0]['User ID'], '001')
            records, _ = process(root, root / 'result.xlsx', DEFAULT_CONFIG)
            self.assertEqual(len(records), 1)

    def test_csv_delimiters_and_preamble(self):
        for delimiter in [',', ';', '\t', '|']:
            with self.subTest(delimiter=delimiter), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / 'training.csv'
                source.write_text('Training report\n\n' + delimiter.join(['User ID', 'Email', 'Assignment Title']) + '\n' + delimiter.join(['001', 'learner@example.org', 'AML']) + '\n')
                frame = read_file(source)[0][1]
                self.assertEqual(frame.iloc[0]['User ID'], '001')
                self.assertEqual(frame.iloc[0]['Assignment Title'], 'AML')

    def test_declared_delimiter_utf16_and_quoted_value(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'training.csv'
            source.write_text('sep=;\nUser ID;Assignment Title\n001;"AML; refresher"\n', encoding='utf-16')
            frame = read_file(source)[0][1]
            self.assertEqual(frame.iloc[0]['Assignment Title'], 'AML; refresher')

    def test_unrecognised_headers_explain_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'training.csv').write_text('Unusual learner code,Unusual course\n001,AML\n')
            with self.assertRaisesRegex(ValueError, 'Unusual learner code'):
                process(root, root / 'result.xlsx', DEFAULT_CONFIG)

    def test_configured_aliases(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'training.csv'
            source.write_text('Report\nCustom identifier;Custom course\n001;AML\n')
            frame = read_file(source, {'User ID': ['Custom identifier'], 'Assignment Title': ['Custom course']})[0][1]
            self.assertEqual(frame.iloc[0]['Custom identifier'], '001')

    def test_grouped_headers_disambiguate_repeated_fields(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "Training_Progress_Report_2026.xlsx"
            pd.DataFrame([
                ["User Info", None, "Assignment Info", None, None, "User Progress", None, None],
                ["ID", "Full Name", "ID", "Begin Date", "End Date", "Begin Date", "Completed Status", "Result Status"],
                ["001", "Example Learner", "A-1", "01.01.2026", "31.12.2026", "02.01.2026", "Completed Within deadline", "Passed"],
            ]).to_excel(source, index=False, header=False, sheet_name="Training Progress Report")

            frame = read_file(source)[0][1]
            self.assertIn("User Info::ID", frame.columns)
            self.assertIn("Assignment Info::Begin Date", frame.columns)
            self.assertIn("User Progress::Begin Date", frame.columns)

            records, errors = process(root, root / "result.xlsx", DEFAULT_CONFIG)
            self.assertEqual(errors, [])
            self.assertEqual(records.loc[0, "User ID"], "001")
            self.assertEqual(records.loc[0, "Assignment ID"], "A-1")
            self.assertEqual(records.loc[0, "Assignment Begin"], pd.Timestamp("2026-01-01"))

    def test_executive_summary_separates_active_and_closed_cohorts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "training.csv"
            pd.DataFrame([
                {"User ID": "1", "Full Name": "Closed", "Assignment ID": "A", "Assignment Status": "Past Assignment", "Assignment Begin": "2025-01-01", "Assignment End": "2025-12-31", "Completion Status": "Completed", "Result Status": "Passed"},
                {"User ID": "2", "Full Name": "Active", "Assignment ID": "B", "Assignment Status": "Active Assignment", "Assignment Begin": "2026-01-01", "Assignment End": "2026-12-31", "Completion Status": "Not Started", "Result Status": ""},
                {"User ID": "3", "Full Name": "Closed 2026", "Assignment ID": "C", "Assignment Status": "Past Assignment", "Assignment Begin": "2026-01-01", "Assignment End": "2026-08-31", "Completion Status": "Not Started", "Result Status": ""},
            ]).to_csv(source, index=False)
            output = root / "result.xlsx"

            process(root, output, DEFAULT_CONFIG)
            ws = load_workbook(output, data_only=True)["Executive Summary"]

            self.assertEqual(ws["A1"].value, "AML Training Reconciliation — Reaudited")
            self.assertEqual(ws["A9"].value, 2)
            self.assertEqual(ws["C9"].value, 1)
            self.assertEqual(ws["H14"].value, "Closed at snapshot")
            self.assertEqual(ws["H15"].value, "Active at snapshot")
            self.assertEqual(ws["A25"].value, "Population Summary")
            self.assertEqual(ws["A26"].value, "Population")

            course_ws = load_workbook(output, data_only=True)["Course Reconciliation"]
            self.assertEqual(course_ws["A1"].value, "Course & Cohort Reconciliation")
            self.assertEqual(course_ws["A5"].value, "Course-Level Summary")
            self.assertEqual(course_ws["A10"].value, "Assignment Cohort Performance")
            self.assertEqual(course_ws["A11"].value, "Assignment ID")
            self.assertEqual(course_ws["C12"].value, "Closed at snapshot")
            self.assertEqual(course_ws["C13"].value, "Active at snapshot")

            year_ws = load_workbook(output, data_only=True)["Year Movement"]
            self.assertEqual(year_ws["A1"].value, "Year Movement")
            self.assertEqual(year_ws["A5"].value, "Assignment Year")
            self.assertEqual(year_ws["B6"].value, "Closed at snapshot")
            self.assertEqual(year_ws["B7"].value, "Active at snapshot")
            self.assertEqual(year_ws["B8"].value, "Closed at snapshot")
            self.assertEqual(year_ws["A12"].value, "Closed-Cohort Trend")
            self.assertEqual(year_ws["A13"].value, "Year")
            self.assertEqual(year_ws["F15"].value, -1)


if __name__ == '__main__':
    unittest.main()
