#!/usr/bin/env python3
"""Generic, config-driven reconciliation and audit workbook verifier.

Raw transaction-level spreadsheets, normalized master sheets, derived columns,
and management summaries can all drift independently. This tool reloads the raw
files and verifies every configured layer without modifying any workbook.

All workbook-specific details live in JSON: sheet names, header rows, mappings,
derived rules, rollups, and dashboard cells. Use a different config to audit a
different reconciliation job without changing this module.

Examples::

    python recon_verifier.py --config my_recon_config.json
    python recon_verifier.py --config my_recon_config.json --report out.md
    python recon_verifier.py --config config.json --raw-glob '/data/raw_*.xlsx' \
        --recon-path /data/reconciliation.xlsx

Exit status is 0 when every check passes, 1 for reconciliation mismatches, and 2
for an unreadable input or invalid configuration.
"""
from __future__ import annotations

import argparse
import glob
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, time as dtime
from pathlib import Path
from typing import Any, Callable

import pandas as pd
from openpyxl import load_workbook


AGG_FUNCS: dict[str, Callable[[pd.Series], Any]] = {
    "count": lambda s: s.count(),
    "nunique": lambda s: s.nunique(dropna=True),
    "sum": lambda s: pd.to_numeric(s, errors="coerce").sum(),
    "mean": lambda s: pd.to_numeric(s, errors="coerce").mean(),
}


@dataclass
class Result:
    name: str
    passed: bool
    detail: str
    samples: list[dict[str, Any]] = field(default_factory=list)


def _blank(value: Any) -> bool:
    return value is None or (not isinstance(value, str) and pd.isna(value))


def _display(value: Any) -> Any:
    if _blank(value):
        return ""
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return value.isoformat()
    if hasattr(value, "item"):
        value = value.item()
    return value


def _equal(left: Any, right: Any) -> bool:
    if _blank(left) and _blank(right):
        return True
    if isinstance(left, (float, int)) and isinstance(right, (float, int)):
        return round(float(left), 6) == round(float(right), 6)
    return str(left).strip() == str(right).strip()


def _text(value: Any) -> str:
    """Return a normalized comparison string without turning nulls into text."""
    return "" if _blank(value) else str(value).strip().casefold()


def _to_seconds(value: Any) -> float | None:
    """Convert Excel/Python times and HH:MM:SS strings to elapsed seconds."""
    if _blank(value):
        return None
    if isinstance(value, dtime):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, pd.Timedelta):
        return value.total_seconds()
    if isinstance(value, str):
        parts = value.strip().split(":")
        if len(parts) == 3:
            try:
                hours, minutes, seconds = (float(part) for part in parts)
                return hours * 3600 + minutes * 60 + seconds
            except ValueError:
                return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _worksheet_frame(path: Path, sheet: str | int, cfg: dict[str, Any]) -> pd.DataFrame:
    """Read exactly the configured Excel rows, including grouped headers."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
        mode = cfg.get("header_mode", "flat")
        if mode == "grouped":
            group_row, field_row = cfg["group_header_row"], cfg["field_header_row"]
            groups, current = [], ""
            for cell in ws[group_row]:
                if not _blank(cell.value):
                    current = str(cell.value).strip()
                groups.append(current)
            fields = [c.value for c in ws[field_row]]
            width = max((i + 1 for i, value in enumerate(fields) if not _blank(value)), default=0)
            headers = []
            for i, value in enumerate(fields[:width]):
                field_name = str(value).strip() if not _blank(value) else f"Unnamed: {i + 1}"
                headers.append(f"{groups[i]}::{field_name}" if groups[i] else field_name)
        else:
            row = cfg["header_row"]
            values = [c.value for c in ws[row]]
            width = max((i + 1 for i, value in enumerate(values) if not _blank(value)), default=0)
            headers = [str(value).strip() if not _blank(value) else f"Unnamed: {i + 1}" for i, value in enumerate(values[:width])]
        end = cfg.get("data_end_row", ws.max_row)
        rows = [[c.value for c in row[:width]] for row in ws.iter_rows(min_row=cfg["data_start_row"], max_row=end)]
        frame = pd.DataFrame(rows, columns=headers)
        return frame.dropna(how="all").reset_index(drop=True)
    finally:
        wb.close()


def _source_name(path: Path, transform: dict[str, str] | None) -> str:
    name = path.name
    if transform:
        name = re.sub(transform["pattern"], transform.get("repl", ""), name)
    return name


def _metric(frame: pd.DataFrame, spec: dict[str, Any]) -> Any:
    agg = spec["agg"]
    if agg not in AGG_FUNCS:
        raise ValueError(f"Unsupported aggregation {agg!r}; choose one of {', '.join(AGG_FUNCS)}")
    return AGG_FUNCS[agg](frame[spec["col"]])


def apply_derived_check(frame: pd.DataFrame, check: dict[str, Any]) -> pd.Series:
    """Recompute one configured derived column from a master-sheet source column."""
    source = check["source_col"]
    if source not in frame:
        raise KeyError(source)
    values = frame[source]
    kind = check["type"]
    yes, no = check.get("true_value", 1), check.get("false_value", 0)

    if kind == "role_evidence":
        marker = _text(check["generic_marker"])
        delimiter = check.get("delimiter", "\n")

        def has_role_evidence(value: Any) -> bool:
            if _blank(value):
                return False
            roles = [part.strip() for part in str(value).split(delimiter) if part.strip()]
            return any(marker not in _text(role) for role in roles)

        return values.map(lambda value: yes if has_role_evidence(value) else no)
    if kind == "value_in_set":
        matches = {_text(value) for value in check["match_values"]}
        return values.map(lambda value: yes if _text(value) in matches else no)
    if kind == "value_equals":
        match = _text(check["match_value"])
        return values.map(lambda value: yes if _text(value) == match else no)
    if kind == "year_of_date":
        return pd.to_datetime(values, errors="coerce").dt.year
    if kind == "duration_bucket":
        numeric = values.map(_to_seconds)
        mask = numeric.notna()
        if "min_seconds" in check:
            mask &= numeric >= check["min_seconds"]
        if "max_seconds" in check:
            mask &= numeric < check["max_seconds"]
        return mask.map(lambda matched: yes if matched else no)
    raise ValueError(f"Unsupported derived check type: {kind}")


def apply_derived_checks(frame: pd.DataFrame, checks: list[dict[str, Any]]) -> dict[str, pd.Series]:
    """Public helper for recomputing all configured derived targets."""
    return {check["target_col"]: apply_derived_check(frame, check) for check in checks}


class Verifier:
    def __init__(self, config: dict[str, Any], config_dir: Path = Path.cwd()):
        self.cfg = config
        self.config_dir = config_dir
        self.results: list[Result] = []
        self.raw = pd.DataFrame()
        self.master = pd.DataFrame()

    def add(self, name: str, passed: bool, detail: str, samples: list[dict[str, Any]] | None = None) -> None:
        self.results.append(Result(name, passed, detail, samples or []))

    def load(self) -> None:
        raw_cfg, recon_cfg = self.cfg["raw"], self.cfg["recon"]
        if raw_cfg.get("header_mode", "flat") not in {"flat", "grouped"}:
            raise ValueError("raw.header_mode must be 'flat' or 'grouped'")
        pattern = Path(raw_cfg["glob"])
        if not pattern.is_absolute():
            pattern = self.config_dir / pattern
        paths = [Path(p) for p in sorted(glob.glob(str(pattern), recursive=True)) if not Path(p).name.startswith("~$")]
        if not paths:
            raise ValueError(f"Raw glob matched no files: {pattern}")
        frames = []
        for path in paths:
            frame = _worksheet_frame(path, raw_cfg.get("sheet_name", 0), raw_cfg)
            frame["__source_file__"] = _source_name(path, recon_cfg.get("source_file_transform"))
            frame["__source_row__"] = range(raw_cfg["data_start_row"], raw_cfg["data_start_row"] + len(frame))
            frames.append(frame)
        self.raw = pd.concat(frames, ignore_index=True)
        recon_path = Path(recon_cfg["path"])
        if not recon_path.is_absolute():
            recon_path = self.config_dir / recon_path
        self.recon_path = recon_path
        self.master = _worksheet_frame(recon_path, recon_cfg["records_sheet"], recon_cfg)

    def row_counts(self) -> None:
        source_col = self.cfg["recon"].get("source_file_column")
        self.add("Total row count", len(self.raw) == len(self.master), f"raw={len(self.raw)}, recon={len(self.master)}")
        if not source_col:
            return
        if source_col not in self.master:
            self.add("Rows per source file", False, f"Recon column {source_col!r} is missing")
            return
        raw_counts = self.raw["__source_file__"].value_counts().to_dict()
        master_counts = self.master[source_col].map(lambda x: "" if _blank(x) else str(x).strip()).value_counts().to_dict()
        names = sorted(set(raw_counts) | set(master_counts))
        bad = [{"source": n, "raw": raw_counts.get(n, 0), "recon": master_counts.get(n, 0)} for n in names if raw_counts.get(n, 0) != master_counts.get(n, 0)]
        self.add("Rows per source file", not bad, f"compared {len(names)} source file(s)", bad[:20])

    def align(self) -> pd.DataFrame:
        join = self.cfg.get("join", {"mode": "positional"})
        raw, master = self.raw.copy(), self.master.copy()
        if join["mode"] == "keys":
            rk, mk = join["raw_keys"], join["master_keys"]
            for column in rk:
                if column not in raw:
                    raise ValueError(f"Raw join column is missing: {column}")
            for column in mk:
                if column not in master:
                    raise ValueError(f"Recon join column is missing: {column}")
            raw["__key__"] = list(map(tuple, raw[rk].fillna("").astype(str).values))
            master["__key__"] = list(map(tuple, master[mk].fillna("").astype(str).values))
            raw["__occ__"] = raw.groupby("__key__").cumcount()
            master["__occ__"] = master.groupby("__key__").cumcount()
            raw = raw.rename(columns={c: c + "__raw" for c in raw if c not in {"__key__", "__occ__"}})
            master = master.rename(columns={c: c + "__master" for c in master if c not in {"__key__", "__occ__"}})
            aligned = raw.merge(master, on=["__key__", "__occ__"], how="outer", indicator=True)
            duplicate_raw = int((raw["__occ__"] > 0).sum())
            duplicate_master = int((master["__occ__"] > 0).sum())
            unmatched = aligned[aligned["_merge"] != "both"]
            samples = [
                {"key": key, "occurrence": occurrence + 1, "side": str(side)}
                for key, occurrence, side in unmatched[["__key__", "__occ__", "_merge"]]
                .head(20)
                .itertuples(index=False, name=None)
            ]
            if duplicate_raw or duplicate_master:
                samples.insert(0, {
                    "problem": "configured keys are not unique",
                    "raw duplicate rows": duplicate_raw,
                    "recon duplicate rows": duplicate_master,
                })
            self.add(
                "Join integrity",
                unmatched.empty and not duplicate_raw and not duplicate_master,
                f"unmatched={len(unmatched)}, raw duplicates={duplicate_raw}, recon duplicates={duplicate_master}",
                samples[:20],
            )
            return aligned
        if join["mode"] != "positional":
            raise ValueError("join.mode must be 'positional' or 'keys'")
        source_col = self.cfg["recon"].get("source_file_column")
        if source_col and source_col in master:
            raw["__position__"] = raw.groupby("__source_file__", dropna=False).cumcount()
            master["__master_source__"] = master[source_col].map(
                lambda value: "" if _blank(value) else str(value).strip()
            )
            master["__position__"] = master.groupby("__master_source__", dropna=False).cumcount()
            raw = raw.rename(columns={c: c + "__raw" for c in raw if c not in {"__source_file__", "__position__"}})
            master = master.rename(columns={c: c + "__master" for c in master if c not in {"__master_source__", "__position__"}})
            aligned = raw.merge(
                master,
                left_on=["__source_file__", "__position__"],
                right_on=["__master_source__", "__position__"],
                how="outer",
                indicator=True,
            )
            unmatched = aligned[aligned["_merge"] != "both"]
            self.add(
                "Join integrity",
                unmatched.empty,
                f"{len(unmatched)} positional row(s) did not pair",
                [
                    {
                        "source": master_source if _blank(raw_source) else raw_source,
                        "position": position + 1,
                        "side": str(side),
                    }
                    for raw_source, position, master_source, side in unmatched[
                        ["__source_file__", "__position__", "__master_source__", "_merge"]
                    ].head(20).itertuples(index=False, name=None)
                ],
            )
            return aligned
        aligned = pd.concat(
            [
                raw.reset_index(drop=True).add_suffix("__raw"),
                master.reset_index(drop=True).add_suffix("__master"),
            ],
            axis=1,
        )
        self.add(
            "Join integrity",
            len(raw) == len(master),
            f"positional rows: raw={len(raw)}, recon={len(master)}",
        )
        return aligned

    def compare_fields(self, aligned: pd.DataFrame) -> None:
        maps = [("Field", self.cfg.get("field_map", {}), False), ("Date field", self.cfg.get("date_field_map", {}), True)]
        date_format = self.cfg.get("raw_date_format")
        for label, mapping, is_date in maps:
            for raw_col, master_col in mapping.items():
                left_name, right_name = raw_col + "__raw", master_col + "__master"
                if left_name not in aligned or right_name not in aligned:
                    missing = raw_col if left_name not in aligned else master_col
                    self.add(f"{label}: {raw_col} → {master_col}", False, f"Column is missing: {missing}")
                    continue
                left, right = aligned[left_name], aligned[right_name]
                displayed_left, displayed_right = left, right
                if is_date:
                    left_blank = left.map(_blank)
                    right_blank = right.map(_blank)
                    parsed_left = pd.to_datetime(left, format=date_format, errors="coerce")
                    parsed_right = pd.to_datetime(right, errors="coerce")
                    invalid = (~left_blank & parsed_left.isna()) | (~right_blank & parsed_right.isna())
                    left, right = parsed_left, parsed_right
                else:
                    invalid = pd.Series(False, index=left.index)
                bad = []
                for index, (a, b) in enumerate(zip(left, right)):
                    if invalid.iloc[index] or not _equal(a, b):
                        bad.append({
                            "row": index + 1,
                            "raw": _display(displayed_left.iloc[index]),
                            "recon": _display(displayed_right.iloc[index]),
                        })
                self.add(f"{label}: {raw_col} → {master_col}", not bad, f"{len(bad)} mismatch(es) across {len(aligned)} row(s)", bad[:20])

    def derived(self) -> None:
        for check in self.cfg.get("derived_checks", []):
            source, target = check["source_col"], check["target_col"]
            result_name = check.get("name", target)
            if source not in self.master or target not in self.master:
                self.add(
                    f"Derived: {result_name}",
                    False,
                    f"Required column missing: {source if source not in self.master else target}",
                )
                continue
            expected = apply_derived_check(self.master, check)
            bad = [{"row": i + 1, "expected": _display(a), "recon": _display(b)} for i, (a, b) in enumerate(zip(expected, self.master[target])) if not _equal(a, b)]
            self.add(f"Derived: {result_name}", not bad, f"{len(bad)} mismatch(es)", bad[:20])

    def summaries(self) -> None:
        for check in self.cfg.get("summary_checks", []):
            summary = _worksheet_frame(self.recon_path, check["sheet_name"], check)
            groups, key_map = check.get("group_by", []), check.get("key_map", {})
            if groups:
                grouped = self.master.groupby(groups, dropna=False)
                expected = grouped.size().reset_index().drop(columns=0)
                for output, metric in check["metrics"].items():
                    expected[output] = grouped[metric["col"]].apply(AGG_FUNCS[metric["agg"]]).values
                expected = expected.rename(columns={key: key_map.get(key, key) for key in groups})
                keys = [key_map.get(key, key) for key in groups]
                compared = expected.merge(summary, on=keys, how="outer", suffixes=("__expected", "__actual"), indicator=True)
            else:
                row = {output: _metric(self.master, metric) for output, metric in check["metrics"].items()}
                compared = pd.concat([pd.DataFrame([row]).add_suffix("__expected"), summary.head(1).add_suffix("__actual")], axis=1)
            bad = []
            for i, row in compared.iterrows():
                for output in check["metrics"]:
                    a, b = row.get(output + "__expected"), row.get(output + "__actual")
                    if not _equal(a, b):
                        bad.append({"row": i + 1, "metric": output, "expected": _display(a), "actual": _display(b)})
            if "_merge" in compared:
                for i, state in compared["_merge"].items():
                    if state != "both":
                        bad.append({"row": i + 1, "problem": str(state)})
            self.add(f"Summary: {check['name']}", not bad, f"{len(bad)} mismatch(es)", bad[:20])

    def cells(self) -> None:
        wb = load_workbook(self.recon_path, read_only=True, data_only=True)
        try:
            for check in self.cfg.get("cell_checks", []):
                frame = self.master.query(check["filter"]) if check.get("filter") else self.master
                expected = _metric(frame, check["metric"])
                actual = wb[check["sheet"]][check["cell"]].value
                self.add(f"Cell: {check['name']}", _equal(expected, actual), f"expected={_display(expected)}, actual={_display(actual)}", [] if _equal(expected, actual) else [{"cell": check["cell"], "expected": _display(expected), "actual": _display(actual)}])
        finally:
            wb.close()

    def run(self) -> list[Result]:
        self.load()
        self.row_counts()
        aligned = self.align()
        self.compare_fields(aligned)
        self.derived()
        self.summaries()
        self.cells()
        return self.results


def write_report(path: Path, results: list[Result]) -> None:
    lines = ["# Reconciliation verification report", "", f"**Overall: {'PASS' if all(r.passed for r in results) else 'FAIL'}**", ""]
    for result in results:
        lines += [f"## {'PASS' if result.passed else 'FAIL'} — {result.name}", "", result.detail, ""]
        if result.samples:
            columns = list(dict.fromkeys(key for sample in result.samples for key in sample))

            def escape(value: Any) -> str:
                return str(_display(value)).replace("|", "\\|").replace("\n", " ")

            lines += ["| " + " | ".join(columns) + " |", "| " + " | ".join("---" for _ in columns) + " |"]
            lines += ["| " + " | ".join(escape(sample.get(column, "")) for column in columns) + " |" for sample in result.samples]
            lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--config", required=True, type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument(
        "--raw-glob",
        help="Override raw.glob (for example, '/data/exports_*.xlsx').",
    )
    parser.add_argument(
        "--recon-path",
        help="Override recon.path with the workbook to verify.",
    )
    parser.add_argument(
        "--data-dir",
        type=Path,
        help="Resolve relative raw/recon paths under this directory.",
    )
    args = parser.parse_args(argv)
    try:
        with args.config.open(encoding="utf-8") as handle:
            config = json.load(handle)
        config_dir = args.data_dir.resolve() if args.data_dir else args.config.resolve().parent
        if args.raw_glob:
            config["raw"]["glob"] = args.raw_glob
        if args.recon_path:
            config["recon"]["path"] = args.recon_path

        results = Verifier(config, config_dir).run()
        for result in results:
            print(f"[{'PASS' if result.passed else 'FAIL'}] {result.name}: {result.detail}")
        if args.report:
            write_report(args.report, results)
        return 0 if all(r.passed for r in results) else 1
    except (OSError, KeyError, ValueError, json.JSONDecodeError) as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
