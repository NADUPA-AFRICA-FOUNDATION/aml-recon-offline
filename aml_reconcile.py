#!/usr/bin/env python3
"""Offline AML/CFT training reconciliation.

Usage:
  python aml_reconcile.py --input ./input \
      --output ./output/AML_Training_Full_Reconciliation.xlsx \
      --config ./config.json

Reads .xlsx/.xlsm/.xls/.csv files recursively. No network/API/AI required.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

DEFAULT_CONFIG: Dict[str, Any] = {
    "report_snapshot_date": "2026-09-01",
    "internal_email_domains": ["safaricom.co.ke"],
    "internal_role_markers": ["safaricom", "employee", "staff", "internal"],
    "critical_fields": [
        "User ID", "Username", "Full Name", "Assignment ID",
        "Assignment Title", "Completion Status"
    ],
    "column_aliases": {
        "User ID": ["user id", "userid", "user_id", "employee id", "learner id", "person id", "User Info::ID"],
        "Full Name": ["full name", "fullname", "name", "learner name", "participant name", "employee name", "User Info::Full Name"],
        "Username": ["username", "email", "email address", "user email", "login", "login id", "User Info::Username"],
        "Roles": ["roles", "role", "user roles", "job role", "learner role", "User Info::Roles"],
        "Staff Category": ["staff category", "staff type", "employee category", "person type", "User Info::Staff category"],
        "Department": ["department", "dept", "business unit", "function", "User Info::Department"],
        "Region": ["region", "area", "zone", "User Info::Region"],
        "Territory": ["territory", "territory name", "User Info::Territory"],
        "Shop/Team": ["shop/team", "shop", "team", "branch", "location", "User Info::Shop Name/Team"],
        "Learning Item ID": ["learning item id", "learning id", "item id", "course id", "Learning Item Info::ID"],
        "Learning Title": ["learning title", "learning item", "course title", "course name", "item title", "Learning Item Info::Title"],
        "Training Category": ["training category", "category", "learning category"],
        "Assignment Type": ["assignment type", "type", "Assignment Info::Type"],
        "Assignment ID": ["assignment id", "assignment_id", "training assignment id", "enrollment id", "enrolment id", "Assignment Info::ID"],
        "Assignment Title": ["assignment title", "training title", "assigned learning", "assignment name", "Assignment Info::Assignement title"],
        "Assignment Status": ["assignment status", "status", "enrollment status", "enrolment status", "Assignment Info::Assignement Status"],
        "Assignment Begin": ["assignment begin", "assignment start", "start date", "assigned date", "begin date", "Assignment Info::Begin Date"],
        "Assignment End": ["assignment end", "assignment due", "end date", "due date", "deadline", "Assignment Info::End Date"],
        "Completion Status": ["completion status", "completion_status", "learning status", "training status", "User Progress::Completed Status"],
        "Result Status": ["result status", "result", "pass status", "outcome", "User Progress::Result Status"],
        "Training Hours": ["training hours", "hours", "duration hours", "learning hours", "User Progress::Training Hours"],
        "Completion Date": ["completion date", "completed date", "date completed", "completion datetime"]
    },
    "status_rules": {
        "completed": [r"^\s*completed\b", r"\bcomplete[d]?\b"],
        "late": [r"out\s*of\s*deadline", r"outof\s*deadline", r"\blate\b", r"after\s*deadline"],
        "started": [r"started.*not.*completed", r"in\s*progress", r"\bstarted\b"],
        "not_started": [r"not\s*started", r"not\s*commenced"],
        "passed": [r"^\s*passed\s*$", r"\bpass(ed)?\b"]
    }
}

OUTPUT_COLUMNS = [
    "Source File", "Participant Key", "User ID", "Full Name", "Username",
    "Roles", "Staff Category", "Department", "Region", "Territory", "Shop/Team",
    "Learning Item ID", "Learning Title", "Training Category", "Assignment Type",
    "Assignment ID", "Assignment Title", "Assignment Status", "Assignment Begin",
    "Assignment End", "Completion Status", "Result Status", "Training Hours",
    "Population", "Assignment Year", "Completed Flag", "Passed Flag",
    "Started Not Completed Flag", "Not Started Flag", "Completed Late Flag",
    "Missing Critical Fields", "Exact Repeat Review"
]


def clean(v: Any) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    s = str(v).strip()
    if s.lower() in {"nan", "none", "nat"}:
        return ""
    return re.sub(r"\s+", " ", s)


def header_key(v: Any) -> str:
    s = clean(v).lower().replace("_", " ").replace("-", " ")
    s = re.sub(r"[^a-z0-9/ ]+", "", s)
    return re.sub(r"\s+", " ", s).strip()


def deep_merge(base: Dict[str, Any], override: Dict[str, Any]) -> Dict[str, Any]:
    for k, v in override.items():
        if isinstance(v, dict) and isinstance(base.get(k), dict):
            deep_merge(base[k], v)
        else:
            base[k] = v
    return base


def load_config(path: Optional[Path]) -> Dict[str, Any]:
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))
    if path:
        with path.open("r", encoding="utf-8") as f:
            deep_merge(cfg, json.load(f))
    return cfg


def source_files(folder: Path) -> List[Path]:
    allowed = {".xlsx", ".xlsm", ".xls", ".csv"}
    return sorted(
        p for p in folder.rglob("*")
        if p.is_file() and p.suffix.lower() in allowed and not p.name.startswith("~$")
        and p.name != "AML_Training_Full_Reconciliation.xlsx"
    )


def table_from_rows(raw: pd.DataFrame, aliases: Dict[str, Sequence[str]]) -> pd.DataFrame:
    """Find a flat or grouped header near the top without fuzzy matching."""
    raw = raw.dropna(how="all").reset_index(drop=True)
    if raw.empty:
        return pd.DataFrame()
    lookup = {
        header_key(label): target
        for target, labels in aliases.items()
        for label in [target, *labels]
    }
    def score(row: pd.Series) -> int:
        return len({lookup[header_key(value)] for value in row if header_key(value) in lookup})

    scores = [score(row) for _, row in raw.head(50).iterrows()]
    header = 0
    # A single matching word in a title is not enough to identify a later header.
    for index, count in enumerate(scores):
        if count >= 2:
            header = index
            break
    field_values = [clean(value) for value in raw.iloc[header]]
    columns = [value or f"Unnamed: {index}" for index, value in enumerate(field_values)]

    # System exports often use a merged group row immediately above a field row.
    # A flat read turns repeated fields such as ID/Title/Begin Date into duplicate
    # labels, so preserve their group context before assigning DataFrame columns.
    duplicate_fields = len([value for value in field_values if value]) != len(
        {header_key(value) for value in field_values if value}
    )
    if header > 0:
        group_values = [clean(value) for value in raw.iloc[header - 1]]
        looks_merged = 1 < sum(bool(value) for value in group_values) < sum(
            bool(value) for value in field_values
        )
        groups, current = [], ""
        for value in group_values:
            if value:
                current = value
            groups.append(current)
        if (duplicate_fields or looks_merged) and len(
            {header_key(value) for value in groups if value}
        ) >= 2:
            columns = [
                f"{groups[index]}::{field}" if groups[index] and field else field or f"Unnamed: {index}"
                for index, field in enumerate(field_values)
            ]
    result = raw.iloc[header + 1:].copy()
    result.columns = columns
    return result.reset_index(drop=True)


def read_file(path: Path, aliases: Optional[Dict[str, Sequence[str]]] = None) -> List[Tuple[str, pd.DataFrame]]:
    aliases = aliases if aliases is not None else DEFAULT_CONFIG["column_aliases"]
    if path.suffix.lower() == ".csv":
        for enc in ("utf-8-sig", "utf-16", "latin-1"):
            try:
                contents = path.read_text(encoding=enc)
                break
            except UnicodeError:
                continue
        # Evaluate common delimiters against known headings, including files with
        # variable-width report titles and Excel's optional sep= declaration.
        declaration = contents.splitlines()[0] if contents.splitlines() else ""
        delimiters = [",", ";", "\t", "|"]
        if declaration.lower().startswith("sep=") and len(declaration) == 5:
            delimiters = [declaration[-1]]
            contents = contents.partition("\n")[2]
        candidates = []
        for delimiter in delimiters:
            rows = list(csv.reader(io.StringIO(contents), delimiter=delimiter))
            frame = table_from_rows(pd.DataFrame(rows), aliases)
            candidates.append((len(column_map(frame, aliases)), frame))
        return [("CSV", max(candidates, key=lambda candidate: candidate[0])[1])]
    sheets = pd.read_excel(path, sheet_name=None, dtype=object, header=None)
    return [(name, table_from_rows(df, aliases)) for name, df in sheets.items() if not df.empty]


def column_map(df: pd.DataFrame, aliases: Dict[str, Sequence[str]]) -> Dict[str, str]:
    available = {header_key(c): str(c) for c in df.columns}
    out = {}
    for target, vals in aliases.items():
        for alias in [target, *vals]:
            if header_key(alias) in available:
                out[target] = available[header_key(alias)]
                break
    return out


def matches(value: Any, patterns: Sequence[str]) -> bool:
    s = clean(value)
    return bool(s) and any(re.search(p, s, re.I) for p in patterns)


def make_participant_key(uid: Any, username: Any, full_name: Any, unique_row: int) -> str:
    if clean(uid):
        return f"ID:{clean(uid)}"
    if clean(username):
        return f"USER:{clean(username).lower()}"
    # Deliberately never merge on name alone.
    safe = re.sub(r"[^a-z0-9]+", "_", clean(full_name).lower()).strip("_") or "unknown"
    return f"UNRESOLVED:{safe}:{unique_row}"


def population(row: pd.Series, cfg: Dict[str, Any]) -> str:
    username = clean(row.get("Username")).lower()
    for domain in cfg["internal_email_domains"]:
        d = clean(domain).lower().lstrip("@")
        if username.endswith("@" + d):
            return "Internal staff"
    role_text = f"{clean(row.get('Roles'))} {clean(row.get('Staff Category'))}".lower()
    if any(clean(m).lower() in role_text for m in cfg["internal_role_markers"] if clean(m)):
        return "Internal staff"
    return "External partner"


def year_from(begin: Any, end: Any, completion: Any, filename: str) -> Any:
    for v in (begin, end):
        dt = pd.to_datetime(v, errors="coerce")
        if pd.notna(dt):
            return int(dt.year)
    found = re.findall(r"\b(20\d{2})\b", filename)
    if found:
        return int(found[0])
    dt = pd.to_datetime(completion, errors="coerce")
    return int(dt.year) if pd.notna(dt) else ""


def normalize(df: pd.DataFrame, path: Path, sheet: str, cfg: Dict[str, Any], offset: int) -> pd.DataFrame:
    cmap = column_map(df, cfg["column_aliases"])
    raw_fields = list(cfg["column_aliases"].keys())
    out = pd.DataFrame(index=df.index)
    for field in raw_fields:
        out[field] = df[cmap[field]] if field in cmap else ""

    out["Source File"] = path.name if sheet == "CSV" else f"{path.name} | {sheet}"

    text_fields = [x for x in raw_fields if x not in {"Assignment Begin", "Assignment End", "Completion Date", "Training Hours"}]
    for c in text_fields:
        out[c] = out[c].map(clean)

    for c in ("Assignment Begin", "Assignment End", "Completion Date"):
        out[c] = pd.to_datetime(out[c], errors="coerce")
    out["Training Hours"] = pd.to_numeric(out["Training Hours"], errors="coerce")

    out["Participant Key"] = [
        make_participant_key(u, e, n, offset + i + 1)
        for i, (u, e, n) in enumerate(zip(out["User ID"], out["Username"], out["Full Name"]))
    ]
    out["Population"] = out.apply(lambda r: population(r, cfg), axis=1)

    rules = cfg["status_rules"]
    out["Completed Flag"] = out["Completion Status"].map(lambda x: int(matches(x, rules["completed"])))
    out["Passed Flag"] = out["Result Status"].map(lambda x: int(matches(x, rules["passed"])))
    out["Completed Late Flag"] = out["Completion Status"].map(lambda x: int(matches(x, rules["late"])))

    started, not_started = [], []
    for cs, ast, done in zip(out["Completion Status"], out["Assignment Status"], out["Completed Flag"]):
        combined = f"{clean(cs)} {clean(ast)}"
        s = int(not done and matches(combined, rules["started"]))
        n = int(not done and matches(combined, rules["not_started"]))
        if s and n:
            # Explicit 'not started' wins over generic started tokens.
            s = 0
        started.append(s)
        not_started.append(n)
    out["Started Not Completed Flag"] = started
    out["Not Started Flag"] = not_started

    out["Assignment Year"] = [
        year_from(b, e, c, src)
        for b, e, c, src in zip(out["Assignment Begin"], out["Assignment End"], out["Completion Date"], out["Source File"])
    ]

    missing = []
    for _, r in out.iterrows():
        absent = [f for f in cfg["critical_fields"] if clean(r.get(f)) == ""]
        missing.append(", ".join(absent))
    out["Missing Critical Fields"] = missing
    out["Exact Repeat Review"] = ""

    for c in OUTPUT_COLUMNS:
        if c not in out:
            out[c] = ""
    return out[OUTPUT_COLUMNS + ["Completion Date"]]


def flag_repeats(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    cols = ["Participant Key", "Assignment ID", "Assignment Title", "Assignment Begin", "Assignment End", "Completion Status", "Result Status"]
    canon = pd.DataFrame(index=out.index)
    for c in cols:
        if pd.api.types.is_datetime64_any_dtype(out[c]):
            canon[c] = out[c].dt.strftime("%Y-%m-%d").fillna("")
        else:
            canon[c] = out[c].fillna("").map(lambda x: clean(x).lower())
    has_identity = out["Assignment ID"].map(clean).ne("") | out["Assignment Title"].map(clean).ne("")
    resolved = ~out["Participant Key"].str.startswith("UNRESOLVED:", na=False)
    dup = canon.duplicated(cols, keep=False) & has_identity & resolved
    out.loc[dup, "Exact Repeat Review"] = "Potential exact repeated assignment record - verify against source"
    return out


def metrics(g: pd.DataFrame, include_pass_rate: bool = False) -> pd.Series:
    rows = len(g)
    data = {
        "Assignment Rows": rows,
        "Unique Participants": g["Participant Key"].nunique(),
        "Completed": int(g["Completed Flag"].sum()),
        "Completion Rate": float(g["Completed Flag"].sum() / rows) if rows else 0.0,
        "Passed": int(g["Passed Flag"].sum()),
        "Started Not Completed": int(g["Started Not Completed Flag"].sum()),
        "Not Started": int(g["Not Started Flag"].sum()),
        "Completed Late": int(g["Completed Late Flag"].sum()),
    }
    if include_pass_rate:
        data["Pass Rate"] = float(g["Passed Flag"].sum() / rows) if rows else 0.0
    return pd.Series(data)


def course_summary(df: pd.DataFrame, active_mask: Optional[pd.Series] = None) -> pd.DataFrame:
    """Summarise the underlying learning item, independently of its campaigns."""
    active_mask = active_mask if active_mask is not None else pd.Series(False, index=df.index)
    rows = []
    for (item_id, title), group in df.groupby(["Learning Item ID", "Learning Title"], dropna=False):
        active = active_mask.reindex(group.index, fill_value=False)
        closed_group, active_group = group.loc[~active], group.loc[active]
        completed = int(group["Completed Flag"].sum())
        rows.append([
            item_id, title, len(group), group["Participant Key"].nunique(),
            group["Assignment ID"].replace("", np.nan).nunique(), completed,
            completed / len(group) if len(group) else 0, len(closed_group),
            int(closed_group["Completed Flag"].sum()),
            closed_group["Completed Flag"].mean() if len(closed_group) else 0,
            len(active_group), active_group["Completed Flag"].mean() if len(active_group) else 0,
        ])
    return pd.DataFrame(rows, columns=[
        "Learning Item ID", "Learning Title", "Assignment Rows", "Unique Participants",
        "Distinct Assignment Cohorts", "Completed", "Completion Rate", "Closed Rows",
        "Closed Completed", "Closed Completion Rate", "Active Rows", "Active Completion Rate",
    ])


def assignment_cohort_summary(df: pd.DataFrame, active_mask: pd.Series) -> pd.DataFrame:
    """Produce one auditable performance row per assignment campaign/cohort."""
    rows = []
    keys = ["Assignment ID", "Assignment Title", "Assignment Year", "Population"]
    for (assignment_id, title, year, population_name), group in df.groupby(keys, dropna=False):
        active = bool(active_mask.reindex(group.index, fill_value=False).any())
        completed = int(group["Completed Flag"].sum())
        rows.append([
            assignment_id, title, "Active at snapshot" if active else "Closed at snapshot",
            join_unique(group["Assignment Status"]), year, population_name, len(group),
            group["Participant Key"].nunique(), completed,
            completed / len(group) if len(group) else 0, int(group["Passed Flag"].sum()),
            int(group["Started Not Completed Flag"].sum()), int(group["Not Started Flag"].sum()),
            int(group["Completed Late Flag"].sum()),
        ])
    return pd.DataFrame(rows, columns=[
        "Assignment ID", "Assignment Title", "Lifecycle", "Assignment Status Values",
        "Assignment Year", "Population", "Rows", "Unique Participants", "Completed",
        "Completion Rate", "Passed", "Started Not Completed", "Not Started", "Completed Late",
    ])


def year_summary(df: pd.DataFrame, active_mask: Optional[pd.Series] = None) -> pd.DataFrame:
    """Report annual results by lifecycle and population for like-for-like use."""
    active_mask = active_mask if active_mask is not None else pd.Series(False, index=df.index)
    working = df.copy()
    working["Lifecycle"] = np.where(
        active_mask.reindex(working.index, fill_value=False),
        "Active at snapshot", "Closed at snapshot",
    )
    result = working.groupby(
        ["Assignment Year", "Lifecycle", "Population"], dropna=False
    ).apply(metrics).reset_index()
    return result[[
        "Assignment Year", "Lifecycle", "Population", "Assignment Rows",
        "Unique Participants", "Completed", "Completion Rate", "Passed",
        "Started Not Completed", "Not Started", "Completed Late",
    ]]


def closed_cohort_trend(df: pd.DataFrame, active_mask: pd.Series) -> pd.DataFrame:
    """Build a year-on-year trend containing closed assignment cohorts only."""
    rows = []
    closed = df.loc[~active_mask.reindex(df.index, fill_value=False)]
    previous_rate: Optional[float] = None
    for year, group in closed.groupby("Assignment Year", dropna=False):
        completed = int(group["Completed Flag"].sum())
        rate = completed / len(group) if len(group) else 0
        rows.append([
            year, len(group), group["Participant Key"].nunique(), completed, rate,
            "" if previous_rate is None else rate - previous_rate,
        ])
        previous_rate = rate
    return pd.DataFrame(rows, columns=[
        "Year", "Closed Rows", "Participants", "Completed", "Completion Rate", "YoY Change (pp)"
    ])


def join_unique(vals: Iterable[Any]) -> str:
    seen, result = set(), []
    for v in vals:
        s = clean(v)
        if s and s not in seen:
            seen.add(s)
            result.append(s)
    return "; ".join(result)


def history(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for key, g in df.groupby("Participant Key", dropna=False):
        years = sorted({int(float(y)) for y in g["Assignment Year"] if clean(y) and re.fullmatch(r"\d{4}(?:\.0)?", str(y))})
        movement = "Single assignment year" if len(years) <= 1 else " → ".join(map(str, years))
        dates = pd.concat([
            pd.to_datetime(g["Completion Date"], errors="coerce"),
            pd.to_datetime(g["Assignment Begin"], errors="coerce"),
            pd.to_datetime(g["Assignment End"], errors="coerce")
        ]).dropna()
        identity = g["Assignment ID"].map(clean).where(g["Assignment ID"].map(clean).ne(""), g["Assignment Title"].map(clean))
        rows.append({
            "Participant Key": key,
            "User ID(s)": join_unique(g["User ID"]),
            "Full Name(s)": join_unique(g["Full Name"]),
            "Username(s)": join_unique(g["Username"]),
            "Population": join_unique(g["Population"]),
            "Assignment Rows": len(g),
            "Distinct Courses": identity[identity.ne("")].nunique(),
            "Completed": int(g["Completed Flag"].sum()),
            "Passed": int(g["Passed Flag"].sum()),
            "Started Not Completed": int(g["Started Not Completed Flag"].sum()),
            "Not Started": int(g["Not Started Flag"].sum()),
            "Completed Late": int(g["Completed Late Flag"].sum()),
            "Assignment Years": ", ".join(map(str, years)),
            "Movement Classification": movement,
            "First Activity": dates.min() if not dates.empty else pd.NaT,
            "Latest Activity": dates.max() if not dates.empty else pd.NaT,
        })
    return pd.DataFrame(rows)


def exception_rows(df: pd.DataFrame) -> pd.DataFrame:
    mask = df["Missing Critical Fields"].map(clean).ne("") | df["Exact Repeat Review"].map(clean).ne("") | df["Participant Key"].str.startswith("UNRESOLVED:", na=False)
    cols = ["Source File", "Participant Key", "User ID", "Full Name", "Username", "Assignment ID", "Assignment Title", "Completion Status", "Result Status", "Missing Critical Fields", "Exact Repeat Review"]
    return df.loc[mask, cols].copy()


def methodology(cfg: Dict[str, Any]) -> pd.DataFrame:
    rows = [
        ["Unit of analysis", "One participant × one training assignment row.", "Preserve multiple legitimate trainings.", "Do not deduplicate legitimate repeat learners.", "Normalized records", "Assignments and unique people are reported separately."],
        ["Participant key", "User ID first; username/email only if User ID is absent.", "Consolidate individual history.", "Exact identifier matching only.", "User ID / Username", "No fuzzy-name merging."],
        ["Population", "Internal email domains and configured role/category markers identify internal staff.", "Separate staff from partners.", "Others default to External partner.", "Username / Roles / Staff Category", "Review config.json for your organisation."],
        ["Completed", "Completion Status matching configured Completed patterns.", "Measure completion.", "Late completion still counts as completed.", "Completion Status", ""],
        ["Passed", "Result Status matching configured Passed patterns.", "Separate pass from completion.", "Binary Passed Flag.", "Result Status", ""],
        ["Started Not Completed", "Started/in-progress and not completed.", "Identify partial completion.", "Separate binary flag.", "Completion/Assignment Status", ""],
        ["Not Started", "Explicit not-started/not-commenced status and not completed.", "Identify no commencement.", "Separate binary flag.", "Completion/Assignment Status", ""],
        ["Completed Late", "Late/out-of-deadline status pattern.", "Identify deadline breach.", "Still counted as completed.", "Completion Status", ""],
        ["Duplicate review", "Same participant, assignment identity, dates, completion status and result.", "Identify suspicious exact repeats.", "Flag only; never auto-delete.", "Assignment-level fields", "Multi-course and multi-year participation is preserved."],
        ["Missing critical fields", "Configured key fields checked for blanks.", "Surface data-quality risks.", "Retain record and list in Exceptions.", "Normalized records", "Configured in config.json."],
        ["Assignment year", "Begin year, then end year, then year in filename, then completion year.", "Protect cohort-year logic.", "One year per assignment row.", "Dates / filename", "Adjust if your LMS defines cohorts differently."]
    ]
    return pd.DataFrame(rows, columns=["Control", "Rule", "Purpose", "Treatment", "Source", "Notes"])


def validations(all_records: pd.DataFrame, cohorts: pd.DataFrame, years: pd.DataFrame, hist: pd.DataFrame) -> List[str]:
    errors = []
    total = len(all_records)
    if int(cohorts["Rows"].sum()) != total:
        errors.append("Course Reconciliation assignment rows do not equal All Records.")
    if int(years["Assignment Rows"].sum()) != total:
        errors.append("Year Movement assignment rows do not equal All Records.")
    if len(hist) != all_records["Participant Key"].nunique():
        errors.append("Participant History rows do not equal unique Participant Keys.")
    states = all_records["Completed Flag"] + all_records["Started Not Completed Flag"] + all_records["Not Started Flag"]
    if (states > 1).any():
        errors.append("Some rows have conflicting completion-state flags.")
    return errors


def executive_population(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for p, g in df.groupby("Population", dropna=False):
        total = len(g)
        completed = int(g["Completed Flag"].sum())
        evidence = (
            "Organisational/job-role marker" if p == "Internal staff"
            else "No configured internal marker / unconfirmed"
        )
        rows.append([p, total, g["Participant Key"].nunique(), completed, completed / total if total else 0, evidence])
    return pd.DataFrame(rows, columns=["Population", "Assignment Rows", "Participants", "Completed", "Completion Rate", "Evidence Basis"])


def lifecycle_tables(df: pd.DataFrame, cfg: Dict[str, Any]) -> tuple[pd.DataFrame, pd.DataFrame, pd.Series]:
    """Return closed-year and snapshot lifecycle views for the management sheet."""
    snapshot = pd.to_datetime(cfg.get("report_snapshot_date"), errors="coerce")
    if pd.isna(snapshot):
        snapshot = pd.Timestamp.today().normalize()
    statuses = df["Assignment Status"].fillna("").astype(str)
    ends = pd.to_datetime(df["Assignment End"], errors="coerce")
    explicitly_active = statuses.str.contains(r"\b(?:active|current|open)\b", case=False, regex=True)
    explicitly_closed = statuses.str.contains(r"\b(?:archived|past|closed|expired)\b", case=False, regex=True)
    active = explicitly_active | (~explicitly_closed & ends.ge(snapshot))
    closed = ~active

    closed_rows = []
    for year, group in df.loc[closed].groupby("Assignment Year", dropna=False):
        total = len(group)
        completed = int(group["Completed Flag"].sum())
        closed_rows.append([
            year, total, group["Participant Key"].nunique(), completed,
            completed / total if total else 0, int(group["Passed Flag"].sum()),
        ])
    years = pd.DataFrame(closed_rows, columns=[
        "Assignment Year", "Closed Rows", "Unique Participants", "Completed", "Completion Rate", "Passed"
    ])

    lifecycle_rows = []
    for label, mask, interpretation in (
        ("Closed at snapshot", closed, "Eligible for performance comparison"),
        ("Active at snapshot", active, "In-flight; do not treat non-completion as final"),
    ):
        group = df.loc[mask]
        total = len(group)
        completed = int(group["Completed Flag"].sum())
        lifecycle_rows.append([label, total, completed, completed / total if total else 0, interpretation])
    lifecycle = pd.DataFrame(lifecycle_rows, columns=["Lifecycle", "Rows", "Completed", "Completion Rate", "Interpretation"])
    return years, lifecycle, active


def format_workbook(path: Path, validation_errors: List[str]) -> None:
    wb = load_workbook(path)
    dark = PatternFill("solid", fgColor="1F4E3D")
    light = PatternFill("solid", fgColor="D9EAD3")
    warning = PatternFill("solid", fgColor="FCE4D6")
    white_bold = Font(color="FFFFFF", bold=True)

    ws = wb["Executive Summary"]
    ws["A1"] = "AML Training Reconciliation — Reaudited"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = dark
    ws.merge_cells("A1:M1")
    ws.merge_cells("A2:M2")
    ws["A2"].alignment = Alignment(wrap_text=True)

    for s in ["Participant History", "All Records", "Exceptions", "Methodology"]:
        sh = wb[s]
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for c in sh[1]:
            c.fill = dark
            c.font = white_bold
            c.alignment = Alignment(wrap_text=True, vertical="center")
        if sh.max_row >= 2:
            try:
                table = Table(displayName=re.sub(r"[^A-Za-z0-9]", "", s)[:20] + "Tbl", ref=f"A1:{get_column_letter(sh.max_column)}{sh.max_row}")
                table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showFirstColumn=False, showLastColumn=False, showColumnStripes=False)
                sh.add_table(table)
            except Exception:
                pass
        headers = {c.value: c.column for c in sh[1]}
        for h in ("Completion Rate", "Pass Rate"):
            if h in headers:
                for r in range(2, sh.max_row + 1):
                    sh.cell(r, headers[h]).number_format = "0.0%"
        for h in ("Assignment Begin", "Assignment End", "First Activity", "Latest Activity"):
            if h in headers:
                for r in range(2, sh.max_row + 1):
                    sh.cell(r, headers[h]).number_format = "yyyy-mm-dd"
        for col in range(1, sh.max_column + 1):
            h = clean(sh.cell(1, col).value)
            width = 38 if h in {"Source File", "Missing Critical Fields", "Exact Repeat Review", "Rule", "Purpose", "Treatment", "Notes"} else 30 if h in {"Full Name", "Assignment Title", "Learning Title", "Username", "Participant Key"} else min(max(len(h)+3, 12), 22)
            sh.column_dimensions[get_column_letter(col)].width = width
        for row in sh.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)

    course_ws = wb["Course Reconciliation"]
    cohort_title_row = next(
        row for row in range(1, course_ws.max_row + 1)
        if course_ws.cell(row, 1).value == "Assignment Cohort Performance"
    )
    cohort_header_row = cohort_title_row + 1
    course_ws.freeze_panes = f"A{cohort_header_row}"
    course_ws.auto_filter.ref = f"A{cohort_header_row}:N{course_ws.max_row}"
    for row_number in (1, 5, 6, cohort_title_row, cohort_header_row):
        for cell in course_ws[row_number]:
            if cell.value is not None:
                cell.fill = dark if row_number == 1 else light
                cell.font = white_bold if row_number == 1 else Font(bold=True)
    for row in course_ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row_number, end_row, header_row in (
        (7, cohort_title_row - 1, 6),
        (cohort_header_row + 1, course_ws.max_row, cohort_header_row),
    ):
        headers = {cell.value: cell.column for cell in course_ws[header_row]}
        for heading in ("Completion Rate", "Closed Completion Rate", "Active Completion Rate"):
            if heading in headers:
                for current in range(row_number, end_row + 1):
                    if isinstance(course_ws.cell(current, headers[heading]).value, (int, float)):
                        course_ws.cell(current, headers[heading]).number_format = "0.0%"
    for column in range(1, course_ws.max_column + 1):
        heading = clean(course_ws.cell(cohort_header_row, column).value) or clean(course_ws.cell(6, column).value)
        width = 42 if heading in {"Learning Title", "Assignment Title", "Assignment Status Values"} else 23
        course_ws.column_dimensions[get_column_letter(column)].width = width
    course_ws.row_dimensions[2].height = 34

    year_ws = wb["Year Movement"]
    trend_title_row = next(
        row for row in range(1, year_ws.max_row + 1)
        if year_ws.cell(row, 1).value == "Closed-Cohort Trend"
    )
    trend_header_row = trend_title_row + 1
    year_ws.freeze_panes = "A5"
    year_ws.auto_filter.ref = f"A5:K{trend_title_row - 3}"
    for row_number in (1, 5, trend_title_row, trend_header_row):
        for cell in year_ws[row_number]:
            if cell.value is not None:
                cell.fill = dark if row_number == 1 else light
                cell.font = white_bold if row_number == 1 else Font(bold=True)
    for row in year_ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for header_row, first_row, last_row in (
        (5, 6, trend_title_row - 3),
        (trend_header_row, trend_header_row + 1, year_ws.max_row),
    ):
        headers = {cell.value: cell.column for cell in year_ws[header_row]}
        for heading in ("Completion Rate", "YoY Change (pp)"):
            if heading in headers:
                for current in range(first_row, last_row + 1):
                    if isinstance(year_ws.cell(current, headers[heading]).value, (int, float)):
                        year_ws.cell(current, headers[heading]).number_format = "0.0%"
    for column in range(1, year_ws.max_column + 1):
        year_ws.column_dimensions[get_column_letter(column)].width = 24
    year_ws.column_dimensions["A"].width = 28
    year_ws.row_dimensions[2].height = 34

    # Executive summary section headers and rates.
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value in {"Assignment Rows", "Closed Assignment Rows", "Assignment Year", "Population"} or r in {4, 8, 12, 13, 18, 25, 26}:
            for c in ws[r]:
                if c.value is not None:
                    c.fill = light
                    c.font = Font(bold=True)
    rate_columns = set()
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and "Rate" in str(cell.value):
                rate_columns.add((cell.row, cell.column))
    for header_row, column in rate_columns:
        for row_number in range(header_row + 1, ws.max_row + 1):
            value = ws.cell(row_number, column).value
            if isinstance(value, (float, int)):
                ws.cell(row_number, column).number_format = "0.0%"
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in {"A": 28, "B": 24, "C": 22, "D": 22, "E": 20, "F": 24, "G": 4,
                       "H": 23, "I": 12, "J": 14, "K": 18, "L": 48, "M": 3}.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[2].height = 38
    ws.freeze_panes = "A4"
    if validation_errors:
        ws["G1"] = "Validation issues"
        ws["G1"].fill = warning
        ws["G1"].font = Font(bold=True)
        for i, err in enumerate(validation_errors, 2):
            ws.cell(i, 7, err)
        ws.column_dimensions["G"].width = 60

    wb.save(path)


def process(input_dir: Path, output_path: Path, cfg: Dict[str, Any]) -> tuple[pd.DataFrame, List[str]]:
    """Generate the workbook and return records and validation notes for the GUI."""
    files = source_files(input_dir)
    if not files:
        raise FileNotFoundError(f"No Excel/CSV files found in {input_dir.resolve()}")

    frames, offset = [], 0
    import_notes = []
    print("Reading source files...")
    for path in files:
        file_count = 0
        for sheet, df in read_file(path, cfg["column_aliases"]):
            headings = ", ".join(str(column) for column in df.columns[:12])
            import_notes.append(f"{path.name} / {sheet}: {len(df)} rows; headings: {headings or 'none'}")
            n = normalize(df, path, sheet, cfg, offset)
            meaningful = n[["User ID", "Username", "Full Name", "Assignment ID", "Assignment Title", "Completion Status"]].fillna("").astype(str).apply(lambda c: c.str.strip()).ne("").any(axis=1)
            n = n.loc[meaningful].copy()
            if not n.empty:
                frames.append(n)
                offset += len(n)
                file_count += len(n)
        print(f"  {path.name}: {file_count:,} records")

    if not frames:
        raise ValueError(
            "No training records could be recognised. Check that the export contains "
            "a header row and learner or assignment data. Expected headings include "
            "User ID, Email, Full Name, Assignment ID, Assignment Title or Completion Status. "
            "Headers may appear within the first 50 non-empty rows. "
            "If your headings differ, add them to column_aliases in the configuration. "
            "Detected tables: " + "; ".join(import_notes)
        )

    all_records = flag_repeats(pd.concat(frames, ignore_index=True))
    hist = history(all_records)
    exc = exception_rows(all_records)
    meth = methodology(cfg)
    pop = executive_population(all_records)
    closed_years, lifecycle, active_mask = lifecycle_tables(all_records, cfg)
    years = year_summary(all_records, active_mask)
    trend = closed_cohort_trend(all_records, active_mask)
    courses = course_summary(all_records, active_mask)
    cohorts = assignment_cohort_summary(all_records, active_mask)
    errs = validations(all_records, cohorts, years, hist)

    total = len(all_records)
    completed = int(all_records["Completed Flag"].sum())
    closed = all_records.loc[~active_mask]
    active = all_records.loc[active_mask]
    headline = pd.DataFrame([[
        total, all_records["Participant Key"].nunique(),
        all_records.loc[all_records["Population"] == "Internal staff", "Participant Key"].nunique(),
        all_records.loc[all_records["Population"] != "Internal staff", "Participant Key"].nunique(),
        completed, completed / total if total else 0,
    ]], columns=["Assignment Rows", "Unique Participants", "Internal Staff Participants",
                 "External / Unconfirmed", "Completed (All)", "Completion Rate (All)"])
    closed_active = pd.DataFrame([[
        len(closed), closed["Completed Flag"].mean() if len(closed) else 0,
        len(active), active["Completed Flag"].mean() if len(active) else 0,
        int(all_records["Completed Late Flag"].sum()), len(exc),
    ]], columns=["Closed Assignment Rows", "Closed Completion Rate", "Active Assignment Rows",
                 "Active Completion Rate", "Completed Late", "Audit-Flagged Records"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    cohort_title_row = max(10, 9 + len(courses))
    trend_title_row = max(12, 8 + len(years))
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        headline.to_excel(writer, sheet_name="Executive Summary", index=False, startrow=3)
        closed_active.to_excel(writer, sheet_name="Executive Summary", index=False, startrow=7)
        closed_years.to_excel(writer, sheet_name="Executive Summary", index=False, startrow=12, startcol=0)
        lifecycle.to_excel(writer, sheet_name="Executive Summary", index=False, startrow=12, startcol=7)
        pop.to_excel(writer, sheet_name="Executive Summary", index=False, startrow=25)
        courses.to_excel(writer, sheet_name="Course Reconciliation", index=False, startrow=5)
        cohorts.to_excel(writer, sheet_name="Course Reconciliation", index=False, startrow=cohort_title_row)
        years.to_excel(writer, sheet_name="Year Movement", index=False, startrow=4)
        trend.to_excel(writer, sheet_name="Year Movement", index=False, startrow=trend_title_row)
        hist.to_excel(writer, sheet_name="Participant History", index=False)
        all_records[OUTPUT_COLUMNS].to_excel(writer, sheet_name="All Records", index=False)
        exc.to_excel(writer, sheet_name="Exceptions", index=False)
        meth.to_excel(writer, sheet_name="Methodology", index=False)

    # Add the narrative after pandas has created the worksheet.
    wb = load_workbook(output_path)
    ws = wb["Executive Summary"]
    snapshot_text = pd.to_datetime(cfg.get("report_snapshot_date"), errors="coerce")
    snapshot_label = snapshot_text.strftime("%-d %b %Y") if pd.notna(snapshot_text) else "the configured"
    ws["A2"] = (f"Management view separates closed cohorts from assignments still active at the {snapshot_label} "
                "source snapshot. Population logic follows configured organisational and job-role evidence.")
    ws["A12"] = "Closed Cohort Year Comparison"
    ws["H12"] = "Snapshot Lifecycle"
    ws["A18"] = "Audit Highlights"
    highlights = [
        ("Population logic", f"{len(all_records.loc[all_records['Population'] == 'Internal staff']):,} assignment rows have configured internal evidence; {len(all_records.loc[all_records['Population'] != 'Internal staff']):,} remain external/unconfirmed."),
        ("Active-assignment distortion", f"{len(active):,} assignments were still active at the snapshot; their {active['Completed Flag'].mean() if len(active) else 0:.1%} completion rate is shown separately from closed cohorts."),
        ("Closed-cohort performance", f"{int(closed['Completed Flag'].sum()):,} of {len(closed):,} closed assignment rows were completed ({closed['Completed Flag'].mean() if len(closed) else 0:.1%})."),
        ("Duplicate and data-quality review", f"{len(exc):,} records are flagged for missing identifiers, unresolved identity, or potential exact repeats and remain in the detailed audit trail."),
    ]
    for row, (label, note) in enumerate(highlights, 19):
        ws.cell(row, 1, label)
        ws.cell(row, 2, note)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=12)
    ws["A25"] = "Population Summary"
    course_ws = wb["Course Reconciliation"]
    course_ws["A1"] = "Course & Cohort Reconciliation"
    course_ws.merge_cells("A1:N1")
    course_ws["A2"] = ("Separates the underlying learning item (course) from assignment campaigns/cohorts. "
                       "Closed and active cohorts are reported separately.")
    course_ws.merge_cells("A2:N2")
    course_ws["A5"] = "Course-Level Summary"
    course_ws.merge_cells("A5:N5")
    course_ws.cell(cohort_title_row, 1, "Assignment Cohort Performance")
    course_ws.merge_cells(start_row=cohort_title_row, start_column=1, end_row=cohort_title_row, end_column=14)
    year_ws = wb["Year Movement"]
    year_ws["A1"] = "Year Movement"
    year_ws.merge_cells("A1:K1")
    year_ws["A2"] = ("Like-for-like trend reporting is based on closed cohorts. Active assignments are retained "
                     "but shown separately to avoid understating current-year performance.")
    year_ws.merge_cells("A2:K2")
    year_ws.cell(trend_title_row, 1, "Closed-Cohort Trend")
    year_ws.merge_cells(start_row=trend_title_row, start_column=1, end_row=trend_title_row, end_column=11)
    wb.save(output_path)
    format_workbook(output_path, errs)

    print("\nResults")
    print(f"  Assignment rows:     {len(all_records):,}")
    print(f"  Unique participants: {all_records['Participant Key'].nunique():,}")
    print(f"  Completed:           {int(all_records['Completed Flag'].sum()):,}")
    print(f"  Passed:              {int(all_records['Passed Flag'].sum()):,}")
    print(f"  Exceptions:          {len(exc):,}")
    print(f"  Output:              {output_path.resolve()}")
    print("  Validation:          " + ("PASS" if not errs else "REVIEW REQUIRED"))
    if errs:
        for e in errs:
            print("    - " + e)
    return all_records, errs


def build(input_dir: Path, output_path: Path, cfg: Dict[str, Any]) -> List[str]:
    """Keep the command-line entry point's validation-only return contract."""
    _, errs = process(input_dir, output_path, cfg)
    return errs


def main() -> int:
    p = argparse.ArgumentParser(description="Offline AML/CFT training reconciliation")
    p.add_argument("--input", default="./input", help="Folder containing raw Excel/CSV exports")
    p.add_argument("--output", default="./output/AML_Training_Full_Reconciliation.xlsx", help="Output workbook path")
    p.add_argument("--config", default=None, help="Optional JSON config")
    args = p.parse_args()
    try:
        errs = build(Path(args.input), Path(args.output), load_config(Path(args.config) if args.config else None))
        return 1 if errs else 0
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
